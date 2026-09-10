from __future__ import annotations

import re
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "5writing" / "assets" / "CUMCM_2026_论文模板.docx"
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


class CumcmWordTemplateTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        if not TEMPLATE.exists():
            raise unittest.SkipTest(f"template has not been built: {TEMPLATE}")
        cls.document = Document(TEMPLATE)
        cls.text = "\n".join(p.text for p in cls.document.paragraphs)

    def test_a4_and_minimum_margins(self) -> None:
        section = self.document.sections[0]
        self.assertAlmostEqual(section.page_width.cm, 21.0, places=1)
        self.assertAlmostEqual(section.page_height.cm, 29.7, places=1)
        for margin in (
            section.top_margin,
            section.right_margin,
            section.bottom_margin,
            section.left_margin,
        ):
            self.assertGreaterEqual(margin.cm, 2.49)

    def test_required_structure_and_order(self) -> None:
        labels = [
            "{{PAPER_TITLE}}",
            "摘  要",
            "1 问题重述",
            "AI工具使用声明",
            "参考文献",
            "附录 A 支撑材料文件列表",
            "附录 B 完整源程序代码",
        ]
        positions = [self.text.index(label) for label in labels]
        self.assertEqual(positions, sorted(positions))
        self.assertNotRegex(self.text, r"(?m)^\s*目录\s*$")

    def test_page_field_is_centered_in_footer(self) -> None:
        footer = self.document.sections[0].footer
        self.assertTrue(any(p.alignment == 1 for p in footer.paragraphs))
        with zipfile.ZipFile(TEMPLATE) as archive:
            footer_xml = "\n".join(
                archive.read(name).decode("utf-8")
                for name in archive.namelist()
                if re.fullmatch(r"word/footer\d+\.xml", name)
            )
        self.assertIn(" PAGE ", footer_xml)

    def test_title_has_no_inherited_border(self) -> None:
        with zipfile.ZipFile(TEMPLATE) as archive:
            styles = ET.fromstring(archive.read("word/styles.xml"))
            document = ET.fromstring(archive.read("word/document.xml"))
        title_style = styles.find(".//w:style[@w:styleId='Title']", NS)
        self.assertIsNotNone(title_style)
        self.assertIsNone(title_style.find("./w:pPr/w:pBdr", NS))
        first_paragraph = document.find(".//w:body/w:p", NS)
        self.assertIsNotNone(first_paragraph)
        self.assertIsNone(first_paragraph.find("./w:pPr/w:pBdr", NS))

    def test_template_fields_are_deliberate_and_detectable(self) -> None:
        fields = set(re.findall(r"\{\{[A-Z0-9_:-]+\}\}", self.text))
        self.assertGreaterEqual(len(fields), 15)
        self.assertIn("{{AI_DECLARATION}}", fields)
        self.assertIn("{{SOURCE_CODE}}", fields)

    def test_symbol_table_has_fixed_geometry_and_repeating_header(self) -> None:
        with zipfile.ZipFile(TEMPLATE) as archive:
            root = ET.fromstring(archive.read("word/document.xml"))
        table = root.find(".//w:tbl", NS)
        self.assertIsNotNone(table)
        tbl_width = table.find("./w:tblPr/w:tblW", NS)
        self.assertEqual(tbl_width.get(f"{{{NS['w']}}}type"), "dxa")
        self.assertEqual(tbl_width.get(f"{{{NS['w']}}}w"), "9072")
        grid_widths = [
            int(node.get(f"{{{NS['w']}}}w"))
            for node in table.findall("./w:tblGrid/w:gridCol", NS)
        ]
        self.assertEqual(grid_widths, [1500, 5472, 2100])
        self.assertIsNotNone(table.find("./w:tr/w:trPr/w:tblHeader", NS))

    def test_package_is_anonymous_and_has_no_review_residue(self) -> None:
        with zipfile.ZipFile(TEMPLATE) as archive:
            names = archive.namelist()
            xml = "\n".join(
                archive.read(name).decode("utf-8", errors="replace")
                for name in names
                if name.endswith(".xml")
            )
        self.assertNotIn("docProps/custom.xml", names)
        self.assertFalse(any("comments" in name for name in names))
        self.assertNotRegex(xml, r"<w:(?:ins|del)\b")
        for forbidden in ("PRIVATE_WORKSPACE", "PRIVATE_USER", "学校名称", "参赛队员", "指导教师"):
            self.assertNotIn(forbidden, xml)


if __name__ == "__main__":
    unittest.main()
