"""Small real DOCX conversion fixture; XML checks do not substitute for rendered QA."""
from __future__ import annotations

import argparse
import json
from pathlib import Path
import zipfile
from xml.etree import ElementTree as ET


def build(path: Path) -> dict:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import qn
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.8)
    section.left_margin = section.right_margin = Inches(.8)
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'; normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    title = doc.add_paragraph('数学建模文稿转换样例', 'Title')
    for run in title.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    doc.add_paragraph('本样例用于在生成全文前检验公式结构、交叉引用、合并表头和源码排版。通过 XML 检查后，仍须检查实际导出的 PDF 页面。')
    doc.add_heading('原生公式', 1)
    p = doc.add_paragraph()
    p._p.append(parse_xml('''<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">
      <m:sSub><m:e><m:r><m:t>z</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub>
      <m:r><m:t>=</m:t></m:r>
      <m:f><m:num><m:rad><m:radPr><m:degHide m:val="1"/></m:radPr><m:deg/>
      <m:e><m:r><m:t>1+x</m:t></m:r></m:e></m:rad></m:num>
      <m:den><m:r><m:t>2</m:t></m:r></m:den></m:f></m:oMath>'''))
    doc.add_heading('交叉引用与合并表头', 1)
    caption = doc.add_paragraph('表 ')
    begin = OxmlElement('w:bookmarkStart'); begin.set(qn('w:id'), '7'); begin.set(qn('w:name'), 'canary_table')
    end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), '7')
    caption._p.append(begin)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'SEQ Table \\* ARABIC')
    run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = '1'; run.append(text); field.append(run)
    caption._p.append(field); caption._p.append(end); caption.add_run(' 结构检查项目')
    table = doc.add_table(rows=4, cols=3)
    table.style = 'Table Grid'
    table.cell(0, 0).text = '项目'; table.cell(0, 1).merge(table.cell(0, 2)).text = '检查结果'
    table.cell(1, 0).text = '公式'; table.cell(1, 1).text = '分式与根式'; table.cell(1, 2).text = '原生 OMML'
    table.cell(2, 0).text = '引用'; table.cell(2, 1).text = '表编号'; table.cell(2, 2).text = '字段更新'
    table.cell(3, 0).text = '源码'; table.cell(3, 1).text = '等宽字体'; table.cell(3, 2).text = '缩进保持'
    for cell in table.rows[0].cells:
        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'DCE6F1'); cell._tc.get_or_add_tcPr().append(shade)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        item = OxmlElement('w:' + edge)
        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            item.set(qn('w:' + key), value)
        borders.append(item)
    table._tbl.tblPr.append(borders)
    for row in table.rows:
        for cell in row.cells:
            margins = OxmlElement('w:tcMar')
            for edge in ('top', 'bottom', 'left', 'right'):
                item = OxmlElement('w:' + edge); item.set(qn('w:w'), '90'); item.set(qn('w:type'), 'dxa'); margins.append(item)
            cell._tc.get_or_add_tcPr().append(margins)
    header = OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(header)
    p = doc.add_paragraph('交叉引用应显示表 ')
    ref = OxmlElement('w:fldSimple'); ref.set(qn('w:instr'), 'REF canary_table \\h')
    run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = '1'; run.append(text); ref.append(run); p._p.append(ref)
    p.add_run('。')
    doc.add_heading('源码样例', 1)
    for line in ['def normalized_overlap(shared, width):', '    if width <= 0:', '        raise ValueError("positive width required")', '    return max(0.0, shared) / width']:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(line); r.font.name = 'Consolas'; r.font.size = Pt(9)
    doc.core_properties.author = ''; doc.core_properties.last_modified_by = ''
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return inspect(path)


def inspect(path: Path) -> dict:
    with zipfile.ZipFile(path) as package:
        tree = ET.fromstring(package.read('word/document.xml'))
    ns = {'m': 'http://schemas.openxmlformats.org/officeDocument/2006/math',
          'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    checks = {name: bool(tree.findall(xpath, ns)) for name, xpath in {
        'fraction': './/m:f', 'radical': './/m:rad', 'subscript': './/m:sSub',
        'merged_header': './/w:gridSpan', 'repeating_header': './/w:tblHeader',
        'reference_field': './/w:fldSimple', 'bookmark': './/w:bookmarkStart'}.items()}
    checks['code_indentation'] = any((e.text or '').startswith('    return') for e in tree.findall('.//w:t', ns))
    instructions = [e.attrib.get('{'+ns['w']+'}instr', '') for e in tree.findall('.//w:fldSimple', ns)]
    checks['ref_and_sequence'] = any(s.startswith('REF ') for s in instructions) and any(s.startswith('SEQ ') for s in instructions)
    return {'schema_version': '1.0', 'structural_checks': checks,
            'status': 'structural_pass' if all(checks.values()) else 'failed',
            'rendered_visual_qa': 'required'}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', required=True, type=Path)
    args = parser.parse_args()
    print(json.dumps(build(args.output), ensure_ascii=False, indent=2))
