#!/usr/bin/env python3
"""Build the editable CUMCM 2026 Word-first paper template.

The generated DOCX is a starting document, not a submission-ready paper. Every
``{{TOKEN}}`` field must be replaced with real, reviewed project content before
the final CUMCM compliance gate can pass.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


TEMPLATE_VERSION = "2026.1"
PAGE_WIDTH_CM = 21.0
PAGE_HEIGHT_CM = 29.7
MARGIN_CM = 2.5
CONTENT_WIDTH_DXA = 9072
TABLE_INDENT_DXA = 120

BLACK = RGBColor(0x00, 0x00, 0x00)
PLACEHOLDER = RGBColor(0x9C, 0x27, 0x2B)
TABLE_HEADER_FILL = "E7E6E6"


def set_run_fonts(run, *, east_asia: str, latin: str, size: float | None = None) -> None:
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), latin)
    rpr.rFonts.set(qn("w:hAnsi"), latin)
    rpr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)


def set_style_fonts(style, *, east_asia: str, latin: str, size: float) -> None:
    style.font.name = latin
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), latin)
    rpr.rFonts.set(qn("w:hAnsi"), latin)
    rpr.rFonts.set(qn("w:eastAsia"), east_asia)


def remove_paragraph_borders(element) -> None:
    ppr = element.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is not None:
        ppr.remove(borders)


def set_repeat_table_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trpr.append(tbl_header)


def shade_cell(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, *, top=80, start=120, bottom=80, end=120) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    tc_mar = tcpr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tcpr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise ValueError("table column widths must sum to the A4 content width")
    table.autofit = False
    tblpr = table._tbl.tblPr

    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tblw.set(qn("w:type"), "dxa")

    tblind = tblpr.find(qn("w:tblInd"))
    if tblind is None:
        tblind = OxmlElement("w:tblInd")
        tblpr.append(tblind)
    tblind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tblind.set(qn("w:type"), "dxa")

    layout = tblpr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblpr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths, strict=True):
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(width))
            tcw.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_page_number(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)

    run = paragraph.add_run()
    set_run_fonts(run, east_asia="宋体", latin="Times New Roman", size=9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    cached = OxmlElement("w:t")
    cached.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, cached, end):
        run._r.append(element)


def add_update_fields_setting(document) -> None:
    settings = document.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def configure_styles(document) -> None:
    styles = document.styles

    normal = styles["Normal"]
    set_style_fonts(normal, east_asia="宋体", latin="Times New Roman", size=10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    set_style_fonts(title, east_asia="黑体", latin="Times New Roman", size=18)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.keep_with_next = True
    remove_paragraph_borders(title.element)

    for name, size, before, after in (
        ("Heading 1", 14, 12, 6),
        ("Heading 2", 12, 10, 5),
        ("Heading 3", 10.5, 8, 4),
    ):
        style = styles[name]
        set_style_fonts(style, east_asia="黑体", latin="Times New Roman", size=size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    body = styles.add_style("CUMCM 正文", 1)
    body.base_style = normal
    set_style_fonts(body, east_asia="宋体", latin="Times New Roman", size=10.5)
    body.paragraph_format.first_line_indent = Pt(21)
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.paragraph_format.space_before = Pt(0)
    body.paragraph_format.space_after = Pt(0)
    body.paragraph_format.line_spacing = 1.25
    body.paragraph_format.widow_control = True

    abstract = styles.add_style("CUMCM 摘要正文", 1)
    abstract.base_style = normal
    set_style_fonts(abstract, east_asia="宋体", latin="Times New Roman", size=10.5)
    abstract.paragraph_format.first_line_indent = Pt(21)
    abstract.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    abstract.paragraph_format.line_spacing = 1.25
    abstract.paragraph_format.space_after = Pt(0)

    caption = styles["Caption"]
    set_style_fonts(caption, east_asia="宋体", latin="Times New Roman", size=9)
    caption.font.color.rgb = BLACK
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(6)
    caption.paragraph_format.keep_with_next = True

    table_text = styles.add_style("CUMCM 表格文字", 1)
    table_text.base_style = normal
    set_style_fonts(table_text, east_asia="宋体", latin="Times New Roman", size=9)
    table_text.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table_text.paragraph_format.space_before = Pt(0)
    table_text.paragraph_format.space_after = Pt(0)
    table_text.paragraph_format.line_spacing = 1.0

    code = styles.add_style("CUMCM 代码", 1)
    code.base_style = normal
    set_style_fonts(code, east_asia="等线", latin="Consolas", size=8)
    code.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    code.paragraph_format.left_indent = Cm(0.5)
    code.paragraph_format.first_line_indent = Cm(0)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 1.0


def add_placeholder(document, token: str, *, style: str = "CUMCM 正文", prefix: str = ""):
    paragraph = document.add_paragraph(style=style)
    if prefix:
        lead = paragraph.add_run(prefix)
        set_run_fonts(lead, east_asia="宋体", latin="Times New Roman", size=10.5)
        lead.bold = True
    run = paragraph.add_run(f"{{{{{token}}}}}")
    set_run_fonts(
        run,
        east_asia="等线" if style == "CUMCM 代码" else "宋体",
        latin="Consolas" if style == "CUMCM 代码" else "Times New Roman",
        size=8 if style == "CUMCM 代码" else 10.5,
    )
    run.font.color.rgb = PLACEHOLDER
    run.italic = True
    return paragraph


def add_symbol_table(document) -> None:
    table = document.add_table(rows=2, cols=3)
    table.style = "Table Grid"
    headers = ("符号", "含义", "单位")
    values = ("{{SYMBOL_1}}", "{{SYMBOL_MEANING_1}}", "{{SYMBOL_UNIT_1}}")
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.style = "CUMCM 表格文字"
        run = paragraph.add_run(text)
        set_run_fonts(run, east_asia="黑体", latin="Times New Roman", size=9)
        run.bold = True
        shade_cell(cell, TABLE_HEADER_FILL)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for index, text in enumerate(values):
        cell = table.rows[1].cells[index]
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.style = "CUMCM 表格文字"
        run = paragraph.add_run(text)
        set_run_fonts(run, east_asia="宋体", latin="Times New Roman", size=9)
        run.font.color.rgb = PLACEHOLDER
        run.italic = True
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_repeat_table_header(table.rows[0])
    set_table_geometry(table, [1500, 5472, 2100])
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def build_document(output: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Cm(PAGE_WIDTH_CM)
    section.page_height = Cm(PAGE_HEIGHT_CM)
    section.top_margin = Cm(MARGIN_CM)
    section.bottom_margin = Cm(MARGIN_CM)
    section.left_margin = Cm(MARGIN_CM)
    section.right_margin = Cm(MARGIN_CM)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.5)
    section.different_first_page_header_footer = False

    configure_styles(document)
    add_page_number(section)
    add_update_fields_setting(document)

    props = document.core_properties
    props.title = "CUMCM 2026 数学建模竞赛论文模板"
    props.subject = "Word-first editable competition paper template"
    props.creator = ""
    props.last_modified_by = ""
    props.keywords = "CUMCM;数学建模;Word模板"
    props.comments = f"Template version {TEMPLATE_VERSION}; replace all template fields."

    title = document.add_paragraph(style="Title")
    remove_paragraph_borders(title._p)
    title_run = title.add_run("{{PAPER_TITLE}}")
    set_run_fonts(title_run, east_asia="黑体", latin="Times New Roman", size=18)
    title_run.bold = True
    title_run.font.color.rgb = PLACEHOLDER

    abstract_label = document.add_paragraph()
    abstract_label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    abstract_label.paragraph_format.space_after = Pt(6)
    label_run = abstract_label.add_run("摘  要")
    set_run_fonts(label_run, east_asia="黑体", latin="Times New Roman", size=12)
    label_run.bold = True

    add_placeholder(document, "ABSTRACT", style="CUMCM 摘要正文")
    keywords = add_placeholder(document, "KEYWORDS", style="Normal", prefix="关键词：")
    keywords.paragraph_format.first_line_indent = Pt(0)
    keywords.paragraph_format.space_before = Pt(8)
    document.add_page_break()

    for heading, token in (
        ("1 问题重述", "PROBLEM_RESTATEMENT"),
        ("2 问题分析", "PROBLEM_ANALYSIS"),
        ("3 模型假设", "MODEL_ASSUMPTIONS"),
    ):
        document.add_heading(heading, level=1)
        add_placeholder(document, token)

    document.add_heading("4 符号说明", level=1)
    add_symbol_table(document)

    document.add_heading("5 模型的建立与求解", level=1)
    add_placeholder(document, "MODEL_OVERVIEW")
    for heading, token in (
        ("5.1 问题一的模型与求解", "PROBLEM_1_MODEL_AND_SOLUTION"),
        ("5.2 问题二的模型与求解", "PROBLEM_2_MODEL_AND_SOLUTION"),
        ("5.3 问题三的模型与求解", "PROBLEM_3_MODEL_AND_SOLUTION"),
    ):
        document.add_heading(heading, level=2)
        add_placeholder(document, token)

    for heading, token in (
        ("6 模型检验与敏感性分析", "VALIDATION_AND_SENSITIVITY"),
        ("7 模型评价与推广", "MODEL_EVALUATION_AND_EXTENSION"),
        ("8 结论", "CONCLUSIONS"),
    ):
        document.add_heading(heading, level=1)
        add_placeholder(document, token)

    document.add_heading("AI工具使用声明", level=1)
    add_placeholder(document, "AI_DECLARATION")
    document.add_heading("参考文献", level=1)
    references = add_placeholder(document, "REFERENCES", style="Normal")
    references.paragraph_format.first_line_indent = Pt(0)

    document.add_page_break()
    document.add_heading("附录 A 支撑材料文件列表", level=1)
    support = add_placeholder(document, "SUPPORT_FILE_LIST", style="Normal")
    support.paragraph_format.first_line_indent = Pt(0)
    document.add_heading("附录 B 完整源程序代码", level=1)
    add_placeholder(document, "SOURCE_CODE", style="CUMCM 代码")

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    build_document(args.output.resolve())


if __name__ == "__main__":
    main()
